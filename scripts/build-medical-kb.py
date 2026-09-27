#!/usr/bin/env python3
"""Build a versioned local reference corpus from pinned, licensed source snapshots.

No network, model calls or spreadsheet formula execution. Requires openpyxl.
Merged cells retain their original anchor coordinates; no forward fill is guessed.
"""
import hashlib
import json
import re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data/medical-kb'


def sha(value):
    return hashlib.sha256(value).hexdigest()


def extract_workbook(source):
    path = DATA / source['file']
    workbook = openpyxl.load_workbook(path, data_only=False)
    sourcesum = sha(path.read_bytes())
    chunks = []
    for sheet in workbook:
        if not ('.DT.' in sheet.title or sheet.title.startswith(('PNC.S.', 'CHE.'))):
            continue
        merged = {}
        for area in sheet.merged_cells.ranges:
            anchor = sheet.cell(area.min_row, area.min_col)
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    merged[(row, col)] = anchor

        def row_text(row):
            parts = []
            for col in range(1, sheet.max_column + 1):
                cell = sheet.cell(row, col)
                origin = merged.get((row, col), cell)
                if origin.value is None:
                    continue
                # Formulas/errors are not clinical content and must not be silently evaluated.
                if origin.data_type in ('f', 'e'):
                    continue
                value = str(origin.value).strip()
                if value:
                    location = cell.coordinate
                    if origin.coordinate != location:
                        location += f' [merged from {origin.coordinate}]'
                    parts.append(f'{location}: {value}')
            return '\n'.join(parts)

        header = '\n'.join(filter(None, (row_text(row) for row in range(1, 7))))
        category = ''
        for row in range(7, sheet.max_row + 1):
            body = row_text(row)
            if not body:
                continue
            # Category rows such as "Fever" are retained alongside subsequent rules.
            originals = [str(c.value).strip() for c in sheet[row] if c.value is not None and c.data_type not in ('f', 'e')]
            if len(originals) == 1 and not re.match(r'(PNC|CHE)[.A-Z0-9]', originals[0]):
                category = originals[0]
                continue
            if not originals:
                continue
            text = f'Sheet: {sheet.title}\n\nTable context and column headers:\n{header}\n\nSection: {category}\n\nOriginal row {row} (merged cells retain their anchors):\n{body}'
            digest = sha(text.encode())
            # Semantic source IDs plus content hashes keep old citations immutable.
            chunk_id = f"{source['id']}-{sha(sheet.title.encode())[:8]}-r{row}-{digest[:10]}"
            chunks.append({'id': chunk_id, 'source_id': source['id'], 'title': sheet.title,
                           'locator': f'{sheet.title} · row {row}; context rows 1–6',
                           'row': row, 'sheet': sheet.title, 'text': text,
                           'text_sha256': digest, 'source_sha256': sourcesum,
                           'scope': source['scope'], 'excerpt_only': True})
    return chunks


def main():
    config = json.loads((ROOT / 'config/medical-kb/sources.json').read_text())
    sources, chunks = [], []
    for source in config['sources']:
        source = {**source, 'retrieved_at': '2026-09-23', 'language': 'en',
                  'source_sha256': sha((DATA / source['file']).read_bytes()), 'clinically_reviewed': False}
        sources.append(source)
        chunks.extend(extract_workbook(source))
    for path in sorted((DATA / 'excerpts').glob('*.json')):
        item = json.loads(path.read_text())
        body = '\n\n'.join(item['sections'])
        if not body or '' in body:
            raise ValueError(f'Invalid extraction: {path}')
        checksum = sha(body.encode())
        source = {'id': item['id'], 'title': item['title'], 'title_zh': item['title_zh'],
                  'publisher': 'Centers for Disease Control and Prevention (CDC)',
                  'url': item['url'], 'version': item['date'], 'source_type': 'Selected official web guidance excerpts',
                  'scope': item['scope'], 'population': item['population'], 'jurisdiction': 'United States / 美国',
                  'license': item['license'], 'license_url': item['license_url'], 'language': 'en',
                  'retrieved_at': item['retrieved_at'], 'source_sha256': checksum, 'clinically_reviewed': False,
                  'limitations': '仅收录所列章节摘录，不是完整页面、指南或药品知识库；未翻译或改写英文正文。',
                  'acquisition': item['acquisition']}
        sources.append(source)
        chunks.append({'id': f"{item['id']}-{checksum[:10]}", 'source_id': item['id'],
                       'title': item['title_zh'], 'locator': 'Selected sections: ' + ' / '.join(s.split('\n')[0].lstrip('# ') for s in item['sections']),
                       'text': body, 'text_sha256': checksum, 'source_sha256': checksum,
                       'scope': item['scope'], 'topics': item['topics'], 'excerpt_only': True})
    assert len({c['id'] for c in chunks}) == len(chunks)
    for chunk in chunks:
        if len(chunk['text']) > 26000:
            raise ValueError('A row exceeds the bounded reader budget: ' + chunk['id'])
    corpus = {'version': config['version'], 'built_at': '2026-09-23', 'sources': sources, 'chunks': chunks}
    output = DATA / 'corpus.json'
    output.write_text(json.dumps(corpus, ensure_ascii=False, indent=2) + '\n')
    (DATA / 'manifest.json').write_text(json.dumps({'version': config['version'], 'sources': len(sources),
        'chunks': len(chunks), 'corpus_sha256': sha(output.read_bytes()),
        'counts': {s['id']: sum(c['source_id'] == s['id'] for c in chunks) for s in sources}}, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'sources': len(sources), 'chunks': len(chunks), 'max_chunk_chars': max(len(c['text']) for c in chunks)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
